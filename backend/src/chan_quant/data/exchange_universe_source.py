from __future__ import annotations

import warnings
from collections.abc import Callable
from dataclasses import dataclass
from datetime import date, datetime, timedelta, timezone
from io import BytesIO

import httpx
from openpyxl import load_workbook
from pydantic import ValidationError

from chan_quant.domain.market_breadth import BreadthUniverseMember

SHANGHAI_TZ = timezone(timedelta(hours=8))
SSE_UNIVERSE_URL = "https://query.sse.com.cn/sseQuery/commonQuery.do"
SZSE_UNIVERSE_URL = "https://www.szse.cn/api/report/ShowReport"


class ExchangeUniverseSourceError(ValueError):
    """Raised when the live exchange lists cannot establish the market denominator."""

    def __init__(self, code: str, message: str) -> None:
        self.code = code
        self.message = message
        super().__init__(f"{code}: {message}")


@dataclass(frozen=True)
class ExchangeUniverseCapture:
    members: tuple[BreadthUniverseMember, ...]
    universe_id: str
    universe_as_of: date
    observed_at: datetime


def _request(
    client: httpx.Client,
    url: str,
    *,
    params: dict[str, str],
    referer: str,
) -> httpx.Response:
    last_error: httpx.HTTPError | None = None
    for _ in range(3):
        try:
            response = client.get(
                url,
                params=params,
                headers={
                    "Accept": "application/json, application/octet-stream, */*",
                    "Referer": referer,
                    "User-Agent": "ChanQuant/0.2 remote-research",
                },
                timeout=httpx.Timeout(connect=5.0, read=20.0, write=5.0, pool=5.0),
            )
        except httpx.HTTPError as error:
            last_error = error
            continue
        if response.status_code >= 500:
            continue
        if response.status_code != 200:
            raise ExchangeUniverseSourceError(
                "EXCHANGE_UNIVERSE_HTTP_ERROR",
                f"{response.url.host} returned HTTP {response.status_code}",
            )
        return response
    if last_error is not None:
        raise ExchangeUniverseSourceError(
            "EXCHANGE_UNIVERSE_HTTP_ERROR",
            f"exchange universe request failed: {last_error}",
        ) from last_error
    raise ExchangeUniverseSourceError(
        "EXCHANGE_UNIVERSE_HTTP_ERROR",
        "exchange universe request returned repeated server errors",
    )


def _compact_date(raw: object, *, field: str) -> date:
    text = str(raw).strip()
    try:
        return date(int(text[:4]), int(text[4:6]), int(text[6:8]))
    except (ValueError, IndexError) as error:
        raise ExchangeUniverseSourceError(
            "EXCHANGE_UNIVERSE_CONTRACT_INVALID",
            f"{field} is not YYYYMMDD",
        ) from error


def _sse_members(client: httpx.Client, *, stock_type: str) -> list[BreadthUniverseMember]:
    response = _request(
        client,
        SSE_UNIVERSE_URL,
        params={
            "sqlId": "COMMON_SSE_CP_GPJCTPZ_GPLB_GP_L",
            "STOCK_TYPE": stock_type,
            "COMPANY_STATUS": "2,4,5,7,8",
            "type": "inParams",
            "isPagination": "true",
            "pageHelp.cacheSize": "1",
            "pageHelp.beginPage": "1",
            "pageHelp.pageSize": "3000" if stock_type == "1" else "1000",
            "pageHelp.pageNo": "1",
        },
        referer="https://www.sse.com.cn/assortment/stock/list/share/",
    )
    try:
        rows = response.json()["result"]
    except (ValueError, KeyError, TypeError) as error:
        raise ExchangeUniverseSourceError(
            "EXCHANGE_UNIVERSE_CONTRACT_INVALID",
            "SSE response does not contain a result array",
        ) from error
    if not isinstance(rows, list):
        raise ExchangeUniverseSourceError(
            "EXCHANGE_UNIVERSE_CONTRACT_INVALID",
            "SSE result must be an array",
        )

    board = "MAIN" if stock_type == "1" else "STAR"
    members: list[BreadthUniverseMember] = []
    for index, row in enumerate(rows, start=1):
        if not isinstance(row, dict):
            raise ExchangeUniverseSourceError(
                "EXCHANGE_UNIVERSE_CONTRACT_INVALID",
                f"SSE row {index} must be an object",
            )
        try:
            code = str(row["A_STOCK_CODE"]).strip()
            name = str(row["SEC_NAME_CN"]).strip()
            delist_raw = str(row.get("DELIST_DATE", "-")).strip()
            members.append(
                BreadthUniverseMember(
                    symbol=f"{code}.SH",
                    exchange="SH",
                    security_name=name,
                    security_type="A_STOCK",
                    board=board,
                    list_date=_compact_date(row["LIST_DATE"], field=f"SSE row {index} LIST_DATE"),
                    delist_date=(
                        None
                        if not delist_raw or delist_raw == "-"
                        else _compact_date(delist_raw, field=f"SSE row {index} DELIST_DATE")
                    ),
                )
            )
        except (KeyError, ValidationError) as error:
            raise ExchangeUniverseSourceError(
                "EXCHANGE_UNIVERSE_CONTRACT_INVALID",
                f"SSE row {index}: {error}",
            ) from error
    return members


def _szse_members(client: httpx.Client) -> list[BreadthUniverseMember]:
    response = _request(
        client,
        SZSE_UNIVERSE_URL,
        params={"SHOWTYPE": "XLSX", "CATALOGID": "1110", "TABKEY": "tab1"},
        referer="https://www.szse.cn/market/product/stock/list/index.html",
    )
    if not response.content.startswith(b"PK"):
        raise ExchangeUniverseSourceError(
            "EXCHANGE_UNIVERSE_CONTRACT_INVALID",
            "SZSE response is not an XLSX workbook",
        )
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            workbook = load_workbook(BytesIO(response.content), read_only=True, data_only=True)
        sheet = workbook.active
        sheet.reset_dimensions()
        rows = sheet.iter_rows(values_only=True)
        header = next(rows)
        indexes = {str(value).strip(): index for index, value in enumerate(header) if value}
        required = ["板块", "A股代码", "A股简称", "A股上市日期"]
        if any(name not in indexes for name in required):
            raise ExchangeUniverseSourceError(
                "EXCHANGE_UNIVERSE_CONTRACT_INVALID",
                "SZSE workbook is missing required A-share columns",
            )

        members: list[BreadthUniverseMember] = []
        for row_number, row in enumerate(rows, start=2):
            code_raw = row[indexes["A股代码"]] if len(row) > indexes["A股代码"] else None
            if code_raw is None or not str(code_raw).strip():
                continue
            board_raw = str(row[indexes["板块"]]).strip()
            board = {"主板": "MAIN", "创业板": "CHINEXT"}.get(board_raw)
            if board is None:
                raise ExchangeUniverseSourceError(
                    "EXCHANGE_UNIVERSE_CONTRACT_INVALID",
                    f"SZSE row {row_number} has unknown board {board_raw}",
                )
            code = str(code_raw).strip().zfill(6)
            name = str(row[indexes["A股简称"]]).strip()
            list_date_raw = row[indexes["A股上市日期"]]
            try:
                list_date = date.fromisoformat(str(list_date_raw).strip())
                members.append(
                    BreadthUniverseMember(
                        symbol=f"{code}.SZ",
                        exchange="SZ",
                        security_name=name,
                        security_type="A_STOCK",
                        board=board,
                        list_date=list_date,
                    )
                )
            except (ValueError, ValidationError) as error:
                raise ExchangeUniverseSourceError(
                    "EXCHANGE_UNIVERSE_CONTRACT_INVALID",
                    f"SZSE row {row_number}: {error}",
                ) from error
    except ExchangeUniverseSourceError:
        raise
    except Exception as error:
        raise ExchangeUniverseSourceError(
            "EXCHANGE_UNIVERSE_CONTRACT_INVALID",
            f"SZSE workbook could not be parsed: {error}",
        ) from error
    finally:
        if "workbook" in locals():
            workbook.close()
    return members


def fetch_exchange_a_share_universe(
    *,
    client: httpx.Client | None = None,
    clock: Callable[[], datetime] | None = None,
    minimum_members: int = 5000,
) -> ExchangeUniverseCapture:
    if minimum_members <= 0:
        raise ValueError("minimum_members must be positive")
    now = clock or (lambda: datetime.now(SHANGHAI_TZ))
    observed_at = now()
    owns_client = client is None
    http_client = client or httpx.Client(follow_redirects=False)
    try:
        members = [
            *_sse_members(http_client, stock_type="1"),
            *_sse_members(http_client, stock_type="8"),
            *_szse_members(http_client),
        ]
    finally:
        if owns_client:
            http_client.close()

    members.sort(key=lambda member: member.symbol)
    symbols = [member.symbol for member in members]
    if len(symbols) != len(set(symbols)):
        raise ExchangeUniverseSourceError(
            "EXCHANGE_UNIVERSE_DUPLICATE_SYMBOL",
            "official exchange lists contain duplicate symbols",
        )
    if len(members) < minimum_members:
        raise ExchangeUniverseSourceError(
            "EXCHANGE_UNIVERSE_INCOMPLETE",
            f"official exchange lists returned only {len(members)} members",
        )
    observed_date = observed_at.astimezone(SHANGHAI_TZ).date()
    return ExchangeUniverseCapture(
        members=tuple(members),
        universe_id=f"sse-szse-live-observed-{observed_date.isoformat()}",
        universe_as_of=observed_date,
        observed_at=observed_at,
    )
